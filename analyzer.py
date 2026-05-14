import io
import re
from typing import Generator

import anthropic
from docx import Document
from docx.shared import Pt

SYSTEM_PROMPT = """Ты — эксперт по Instagram-маркетингу с 10-летним опытом. Ты помогаешь авторам и брендам развивать аккаунты, создавать вовлекающий контент и выстраивать стратегию присутствия в социальных сетях.

Правила ответа:
- Используй Markdown для форматирования
- Будь конкретным: давай реальные примеры тем, хуков, хэштегов
- Не давай общих советов — только практические действия
- В контент-плане создавай полноценную таблицу на указанное количество дней"""

LANGUAGES = ["Русский", "English", "Қазақша", "Español", "Deutsch"]


def _lang_line(lang: str) -> str:
    return {
        "Русский": "Отвечай строго на русском языке.",
        "English": "Respond strictly in English.",
        "Қазақша": "Жауапты қазақ тілінде жаз.",
        "Español": "Responde estrictamente en español.",
        "Deutsch": "Antworte ausschließlich auf Deutsch.",
    }.get(lang, f"Respond in {lang}.")


def build_profile_prompt(data: dict, plan_days: int, lang: str = "Русский") -> str:
    return f"""{_lang_line(lang)}

Проанализируй Instagram-профиль и составь полный контент-план.

## Данные о профиле

- **Ниша / тематика:** {data['niche']}
- **Количество подписчиков:** {data['followers']:,}
- **Количество постов:** {data['posts_count']}
- **Целевая аудитория:** {data['target_audience'] or 'не указана'}
- **Цель аккаунта:** {data['goal']}
- **Тон общения:** {data['tone']}
- **Текущий контент:** {data['current_content'] or 'не указано'}

## Анализ профиля

**Сильные стороны** (минимум 3 пункта на основе данных)

**Слабые стороны / точки роста** (минимум 3 пункта)

**Целевая аудитория** — уточни портрет: возраст, интересы, боли, мечты

**Рекомендуемый Tone of Voice**

---

## Контент-план на {plan_days} дней

| День | Формат | Тема поста | Хук (первая фраза) | CTA |
|------|--------|-----------|-------------------|-----|

Заполни все {plan_days} строк. Чередуй форматы: Рилс, Карусель, Фото, Сторис-серия, Экспертный пост.

---

## Советы по росту

**Хэштеги** — 15–20 конкретных, группами: нишевые / средние / широкие

**Лучшее время публикаций** — дни и часы

**Форматы с наибольшим охватом** в этой нише

**Быстрые wins** — 3 действия, которые дадут результат уже на этой неделе

---

## Идеи для постов и рилсов

5 конкретных идей:

**Идея N — [Название]**
- Формат:
- Тема:
- Хук:
- Структура:
- CTA:"""


def build_niche_prompt(data: dict, plan_days: int, lang: str = "Русский") -> str:
    return f"""{_lang_line(lang)}

Составь контент-стратегию для Instagram-аккаунта в нише.

- **Ниша:** {data['niche']}
- **Цель:** {data['goal']}

## Анализ ниши

**Возможности** — почему ниша перспективна в Instagram

**Целевая аудитория** — портрет: возраст, интересы, боли, мечты

**Рекомендуемый Tone of Voice**

**Как выделиться среди конкурентов**

---

## Контент-план на {plan_days} дней

| День | Формат | Тема поста | Хук (первая фраза) | CTA |
|------|--------|-----------|-------------------|-----|

Заполни все {plan_days} строк.

---

## Советы по росту

**Хэштеги** — 15–20 конкретных, группами

**Лучшее время публикаций**

**Форматы с наибольшим охватом**

**Быстрые wins для старта с нуля**

---

## Идеи для постов и рилсов

5 конкретных идей:

**Идея N — [Название]**
- Формат:
- Тема:
- Хук:
- Структура:
- CTA:"""


def build_caption_prompt(data: dict, lang: str = "Русский") -> str:
    return f"""{_lang_line(lang)}

Напиши 3 варианта подписи (caption) для Instagram-поста.

- **Тема поста:** {data['topic']}
- **Тон:** {data['tone']}
- **Цель подписи:** {data['goal']}
- **Целевая аудитория:** {data.get('audience') or 'не указана'}

Для каждого варианта используй структуру:

**Вариант N — [стиль]**

[Хук — цепляющая первая фраза]

[Тело — 3–5 абзацев через пустую строку]

[CTA — призыв к действию]

[Хэштеги — 10–15 штук]

---

Сделай варианты разными по стилю: эмоциональный, экспертный, через личную историю."""


def build_bio_prompt(data: dict, lang: str = "Русский") -> str:
    return f"""{_lang_line(lang)}

Оптимизируй биографию Instagram-аккаунта.

- **Текущая Bio:** {data['current_bio']}
- **Ниша / чем занимаешься:** {data['niche']}
- **Целевая аудитория:** {data.get('audience') or 'не указана'}
- **Цель профиля:** {data['goal']}

Дай 3 варианта улучшенной Bio (до 150 символов каждая):

**Вариант N — [стиль]**
```
[текст Bio]
```
Почему это работает: [1–2 предложения]

---

После вариантов:

## Советы по оформлению профиля

- Имя пользователя (username) — рекомендации
- Имя в профиле (для SEO-поиска)
- Аватар
- Актуальные сторис (о чём делать)
- Ссылка в Bio — что поставить"""


def build_competitor_prompt(data: dict, plan_days: int, lang: str = "Русский") -> str:
    niches_list = "\n".join(f"- {n}" for n in data["niches"] if n.strip())
    return f"""{_lang_line(lang)}

Проанализируй конкурентные ниши в Instagram и найди возможности для выделения.

## Ниши для анализа:
{niches_list}

- **Моя основная ниша:** {data['my_niche']}
- **Цель:** {data['goal']}

## Сравнительный анализ ниш

Для каждой ниши:

**[Ниша]**
- Насыщенность рынка:
- Типичный контент:
- Слабые места конкурентов:
- Незанятые темы:

---

## Уникальное позиционирование

**Чем выделиться** — конкретные идеи, которые мало кто делает

**Контент на пересечении ниш** — где пересечение тем даёт преимущество

---

## Контент-план на {plan_days} дней (с учётом анализа конкурентов)

| День | Формат | Тема | Почему выделяет | CTA |
|------|--------|------|----------------|-----|

Заполни {plan_days} строк с темами, которых нет у конкурентов.

---

## Хэштеги с низкой конкуренцией

15–20 хэштегов, где меньше борьбы за внимание"""


def stream_analysis(
    api_key: str,
    prompt: str,
    usage_container: dict | None = None,
) -> Generator[str, None, None]:
    client = anthropic.Anthropic(api_key=api_key)
    with client.messages.stream(
        model="claude-sonnet-4-6",
        max_tokens=8096,
        system=[{"type": "text", "text": SYSTEM_PROMPT, "cache_control": {"type": "ephemeral"}}],
        messages=[{"role": "user", "content": prompt}],
    ) as stream:
        for text in stream.text_stream:
            yield text
        if usage_container is not None:
            msg = stream.get_final_message()
            usage_container["input_tokens"] = msg.usage.input_tokens
            usage_container["output_tokens"] = msg.usage.output_tokens


def stream_chat(
    api_key: str,
    messages: list,
    usage_container: dict | None = None,
) -> Generator[str, None, None]:
    client = anthropic.Anthropic(api_key=api_key)
    with client.messages.stream(
        model="claude-sonnet-4-6",
        max_tokens=4096,
        system=[{"type": "text", "text": SYSTEM_PROMPT, "cache_control": {"type": "ephemeral"}}],
        messages=messages,
    ) as stream:
        for text in stream.text_stream:
            yield text
        if usage_container is not None:
            msg = stream.get_final_message()
            usage_container["input_tokens"] = msg.usage.input_tokens
            usage_container["output_tokens"] = msg.usage.output_tokens


def _add_inline(para, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            para.add_run(part[2:-2]).bold = True
        else:
            para.add_run(part)


def markdown_to_docx(text: str) -> bytes:
    doc = Document()
    lines = text.split("\n")
    i = 0
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        if not table_rows:
            return
        cols = max(len(r) for r in table_rows)
        t = doc.add_table(rows=len(table_rows), cols=cols)
        t.style = "Table Grid"
        for ri, row in enumerate(table_rows):
            for ci, cell in enumerate(row):
                if ci < cols:
                    t.cell(ri, ci).text = cell
                    if ri == 0:
                        for p in t.cell(ri, ci).paragraphs:
                            for run in p.runs:
                                run.bold = True
        table_rows.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if s.startswith("|") and s.endswith("|"):
            if re.match(r"^\|[-| :]+\|$", s):
                i += 1
                continue
            table_rows.append([c.strip() for c in s.split("|")[1:-1]])
            i += 1
            continue

        flush_table()

        if s.startswith("## "):
            doc.add_heading(s[3:], level=1)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=2)
        elif s.startswith("#### "):
            doc.add_heading(s[5:], level=3)
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, s[2:])
        elif s == "---":
            doc.add_paragraph("─" * 60)
        elif s.startswith("```"):
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph("\n".join(code))
            for run in p.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(9)
        elif s:
            p = doc.add_paragraph()
            _add_inline(p, s)

        i += 1

    flush_table()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
